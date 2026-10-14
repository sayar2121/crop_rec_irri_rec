import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_full_project_doc():
    readme_path = "/Users/sayarpaul/Project/crop_rec_irri_rec/README.md"
    doc_path = "/Users/sayarpaul/Project/crop_rec_irri_rec/DASP_DAFO_Research_Documentation.docx"
    
    with open(readme_path, "r", encoding="utf-8") as f:
        text = f.read()

    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Style definitions
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)

    lines = text.split("\n")
    i = 0
    in_code_block = False
    code_lines = []
    
    while i < len(lines):
        line = lines[i]
        
        # Check code fence
        if line.strip().startswith("```"):
            if in_code_block:
                # Flush code block
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                set_cell_shading(cell, "F1F5F9")
                set_cell_margins(cell, 120, 120, 160, 160)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run("\n".join(code_lines))
                run.font.name = 'Courier New'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(15, 23, 42)
                doc.add_paragraph() # spacer
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        stripped = line.strip()
        
        # Skip empty lines
        if not stripped:
            i += 1
            continue

        # Check Table
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
                
            # Parse table
            if len(table_lines) >= 2:
                # filter out separator line
                parsed_rows = []
                for tline in table_lines:
                    cols = [c.strip() for c in tline.strip("|").split("|")]
                    if all(re.match(r'^:?-+:?$', c) for c in cols if c):
                        continue # separator
                    parsed_rows.append(cols)
                    
                if parsed_rows:
                    num_cols = max(len(r) for r in parsed_rows)
                    doc_table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
                    doc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    
                    for row_idx, r_data in enumerate(parsed_rows):
                        for col_idx in range(num_cols):
                            val = r_data[col_idx] if col_idx < len(r_data) else ""
                            cell = doc_table.cell(row_idx, col_idx)
                            set_cell_margins(cell, 80, 80, 120, 120)
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            run = p.add_run(val)
                            
                            if row_idx == 0:
                                set_cell_shading(cell, "0F172A") # dark slate
                                run.bold = True
                                run.font.size = Pt(9.5)
                                run.font.color.rgb = RGBColor(255, 255, 255)
                            else:
                                if row_idx % 2 == 1:
                                    set_cell_shading(cell, "F8FAFC")
                                else:
                                    set_cell_shading(cell, "FFFFFF")
                                run.font.size = Pt(9.0)
                                run.font.color.rgb = RGBColor(30, 41, 59)
                    doc.add_paragraph()
            continue

        # Check Image: ![Alt text](path)
        img_match = re.search(r'!\[(.*?)\]\((.*?)\)', stripped)
        if img_match:
            caption = img_match.group(1)
            img_path = img_match.group(2)
            
            # Resolve relative/absolute path
            if not os.path.isabs(img_path):
                img_path = os.path.join("/Users/sayarpaul/Project/crop_rec_irri_rec", img_path)
                
            if os.path.exists(img_path):
                try:
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run()
                    run.add_picture(img_path, width=Inches(5.8))
                    
                    if caption:
                        cp = doc.add_paragraph()
                        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        crun = cp.add_run(f"Figure: {caption}")
                        crun.italic = True
                        crun.font.size = Pt(9.5)
                        crun.font.color.rgb = RGBColor(100, 116, 139)
                except Exception as e:
                    print(f"Error adding image {img_path}: {e}")
            i += 1
            continue

        # Check Headings
        if stripped.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(stripped[2:])
            run.bold = True
            run.font.size = Pt(22)
            run.font.color.rgb = RGBColor(15, 23, 42)
            i += 1
            continue
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(stripped[3:])
            run.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(30, 58, 138) # Deep navy/blue
            i += 1
            continue
        elif stripped.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[4:])
            run.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(13, 148, 136) # Teal
            i += 1
            continue
        elif stripped.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(stripped[5:])
            run.bold = True
            run.font.size = Pt(11.5)
            run.font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue

        # Check Blockquote
        if stripped.startswith("> "):
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            set_cell_shading(cell, "EFF6FF") # Soft ice blue
            set_cell_margins(cell, 80, 80, 140, 140)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(stripped[2:])
            run.italic = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(29, 78, 216)
            i += 1
            continue

        # Check Horizontal Rule
        if stripped in ["---", "***", "___"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            # Add subtle divider line
            run = p.add_run("_________________________________________________________________________________")
            run.font.color.rgb = RGBColor(203, 213, 225)
            run.font.size = Pt(8)
            i += 1
            continue

        # Check Bullet list
        if stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            bullet_text = stripped[2:]
            _parse_inline_formatting(p, bullet_text)
            i += 1
            continue

        # Check Numbered list
        num_match = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if num_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            _parse_inline_formatting(p, num_match.group(2))
            i += 1
            continue

        # Standard paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        _parse_inline_formatting(p, stripped)
        i += 1

    doc.save(doc_path)
    print(f"Successfully generated DOCX at {doc_path} (size: {os.path.getsize(doc_path)} bytes)")

def _parse_inline_formatting(paragraph, text):
    # Handles bold **text**, inline code `code`, and plain text
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(180, 83, 9)
        else:
            paragraph.add_run(token)

if __name__ == "__main__":
    create_full_project_doc()
